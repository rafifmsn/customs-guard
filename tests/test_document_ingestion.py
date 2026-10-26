"""Automated tests for CustomsGuard Multimodal Document Ingestion."""

import io
import json
from unittest.mock import patch
import docx
import pytest
from fastapi.testclient import TestClient

from api.main import app

client = TestClient(app)


def test_ingest_json_file_success():
    """Verify JSON commercial invoice upload parses directly and validates."""
    sample_json = {
        "shipment_id": "SHP-JSON-TEST-001",
        "invoice_number": "INV-2026-9999",
        "importer_name": "PT Test Importer",
        "items": [
            {
                "item_id": 1,
                "declared_description": "Smartphones 256GB",
                "declared_hs_code": "8517.13",
                "quantity": 10,
                "unit_price": 500.0,
                "total_value": 5000.0,
                "declared_duty_rate": 0.0,
            }
        ],
    }
    file_bytes = json.dumps(sample_json).encode("utf-8")
    resp = client.post(
        "/api/v1/ingest/document",
        files={"file": ("invoice.json", file_bytes, "application/json")},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "complete"
    assert data["invoice_data"]["shipment_id"] == "SHP-JSON-TEST-001"
    assert len(data["invoice_data"]["items"]) == 1
    assert data["invoice_data"]["items"][0]["declared_hs_code"] == "8517.13"


def test_ingest_empty_file_fails():
    """Verify empty document upload returns failed status."""
    resp = client.post(
        "/api/v1/ingest/document",
        files={"file": ("empty.pdf", b"", "application/pdf")},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "failed"
    assert "empty" in data["error_message"].lower()


def test_ingest_docx_file_success():
    """Verify DOCX invoice extracts line items and populates InvoicePayload."""
    doc = docx.Document()
    doc.add_heading("Commercial Invoice", level=1)
    doc.add_paragraph("Shipment: SHP-DOCX-101")
    table = doc.add_table(rows=1, cols=4)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Item"
    hdr_cells[1].text = "HS Code"
    hdr_cells[2].text = "Description"
    hdr_cells[3].text = "Total Value"

    row_cells = table.add_row().cells
    row_cells[0].text = "1"
    row_cells[1].text = "8471.30"
    row_cells[2].text = "Portable Laptops"
    row_cells[3].text = "$12,500.00"

    buf = io.BytesIO()
    doc.save(buf)
    file_bytes = buf.getvalue()

    resp = client.post(
        "/api/v1/ingest/document",
        files={"file": ("invoice.docx", file_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "complete"
    items = data["invoice_data"]["items"]
    assert len(items) >= 1
    assert items[0]["declared_hs_code"] == "8471.30"


def test_ingest_pdf_text_file_success():
    """Verify text-based PDF invoice extracts line items cleanly."""
    import pypdf

    # Build simple in-memory PDF with pypdf
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    writer.write(buf)
    file_bytes = buf.getvalue()

    # When PDF contains text or mocked Langflow agent output
    mock_agent_output = json.dumps(
        {
            "status": "complete",
            "invoice_data": {
                "shipment_id": "SHP-PDF-EXTRACTED-01",
                "items": [
                    {
                        "item_id": 1,
                        "declared_description": "Ultrasonic Scanner System",
                        "declared_hs_code": "9018.12",
                        "quantity": 2,
                        "unit_price": 25000.0,
                        "total_value": 50000.0,
                        "declared_duty_rate": 0.0,
                    }
                ],
            },
        }
    )

    with patch("api.main._try_langflow_ingest_flow", return_value=mock_agent_output), patch(
        "api.main._extract_text_from_file_bytes", return_value="Ultrasonic Scanner System HS: 9018.12 Value: $50000.00"
    ):
        resp = client.post(
            "/api/v1/ingest/document",
            files={"file": ("invoice.pdf", file_bytes, "application/pdf")},
        )

    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "complete"
    assert data["invoice_data"]["shipment_id"] == "SHP-PDF-EXTRACTED-01"
    assert data["invoice_data"]["items"][0]["declared_hs_code"] == "9018.12"


def test_ingest_unreadable_or_missing_hs_rejected():
    """Verify document lacking HS codes or valid commodities is rejected with clear error."""
    plain_text = "This is a random greeting letter with no customs or invoice data."
    resp = client.post(
        "/api/v1/ingest/document",
        files={"file": ("memo.txt", plain_text.encode("utf-8"), "text/plain")},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "rejected"
    assert "unable to extract" in data["error_message"].lower() or "missing" in data["error_message"].lower()


def test_ingest_handles_conversational_llm_wrappers():
    """Verify agent output with conversational text and markdown blocks is cleaned and validated."""
    messy_llm_text = """Here is the verified extraction:
```json
{
  "status": "complete",
  "invoice_data": {
    "shipment_id": "SHP-CLEANED-99",
    "importer_name": "UNSPECIFIED",
    "items": [
      {
        "item_id": 1,
        "declared_description": "Bluetooth Transmitters",
        "declared_hs_code": "8517.62",
        "quantity": 100,
        "unit_price": 25.0,
        "total_value": 2500.0,
        "declared_duty_rate": 0.0
      }
    ]
  }
}
```
Let me know if you need anything else!"""

    with patch("api.main._try_langflow_ingest_flow", return_value=messy_llm_text), patch(
        "api.main._extract_text_from_file_bytes", return_value="Bluetooth Transmitters 8517.62"
    ):
        resp = client.post(
            "/api/v1/ingest/document",
            files={"file": ("invoice.pdf", b"fake_pdf_bytes", "application/pdf")},
        )

    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "complete"
    assert data["invoice_data"]["shipment_id"] == "SHP-CLEANED-99"
    assert data["invoice_data"]["items"][0]["declared_hs_code"] == "8517.62"

